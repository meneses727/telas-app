"""Propuesta de compra (pantalla 10.7): el documento final editable que el
planificador exporta para gestionar la compra real con el proveedor.
"""
import math
from io import BytesIO

from sqlalchemy.orm import Session, joinedload

from app.models.recomendacion import PropuestaCompra, RecomendacionCompra
from app.models.tela import Tela


def _metros_finales(item: RecomendacionCompra) -> float | None:
    if item.metros_ajustados is not None:
        return float(item.metros_ajustados)
    if item.metros_sugeridos is not None:
        return float(item.metros_sugeridos)
    return None


def obtener_detalle(db: Session, propuesta_id: int) -> dict:
    propuesta = (
        db.query(PropuestaCompra)
        .options(joinedload(PropuestaCompra.usuario), joinedload(PropuestaCompra.items).joinedload(RecomendacionCompra.tela))
        .filter(PropuestaCompra.id == propuesta_id)
        .first()
    )
    if propuesta is None:
        raise ValueError("Propuesta no encontrada")

    filas = []
    total_metros = 0.0
    total_rollos = 0
    for item in propuesta.items:
        metros = _metros_finales(item)
        filas.append(
            {
                "item_id": item.id,
                "tela_id": item.tela_id,
                "tela_codigo": item.tela.codigo,
                "tela_nombre": item.tela.nombre,
                "pct_reposicion": float(item.pct_reposicion) if item.pct_reposicion is not None else None,
                "metros_sugeridos": float(item.metros_sugeridos) if item.metros_sugeridos is not None else None,
                "metros_ajustados": float(item.metros_ajustados) if item.metros_ajustados is not None else None,
                "metros_finales": metros,
                "rollos": item.rollos,
                "no_concluyente": item.no_concluyente,
                "motivo_ajuste": item.motivo_ajuste,
            }
        )
        if metros is not None:
            total_metros += metros
        if item.rollos:
            total_rollos += item.rollos

    return {
        "propuesta_id": propuesta.id,
        "estado": propuesta.estado,
        "fecha": propuesta.fecha,
        "responsable": propuesta.usuario.nombre if propuesta.usuario else None,
        "items": filas,
        "total_metros": round(total_metros, 1),
        "total_rollos": total_rollos,
    }


def actualizar_item(
    db: Session, propuesta_id: int, item_id: int, metros_ajustados: float, motivo_ajuste: str | None
) -> dict:
    item = (
        db.query(RecomendacionCompra)
        .filter(RecomendacionCompra.id == item_id, RecomendacionCompra.propuesta_id == propuesta_id)
        .first()
    )
    if item is None:
        raise ValueError("Ítem no encontrado en esta propuesta")

    tela = db.query(Tela).filter(Tela.id == item.tela_id).first()
    item.metros_ajustados = metros_ajustados
    item.motivo_ajuste = motivo_ajuste
    item.rollos = math.ceil(metros_ajustados / float(tela.m_por_rollo)) if metros_ajustados > 0 else 0
    db.commit()
    return obtener_detalle(db, propuesta_id)


def eliminar_item(db: Session, propuesta_id: int, item_id: int) -> dict:
    item = (
        db.query(RecomendacionCompra)
        .filter(RecomendacionCompra.id == item_id, RecomendacionCompra.propuesta_id == propuesta_id)
        .first()
    )
    if item is None:
        raise ValueError("Ítem no encontrado en esta propuesta")
    db.delete(item)
    db.commit()
    return obtener_detalle(db, propuesta_id)


def cerrar_propuesta(db: Session, propuesta_id: int) -> dict:
    propuesta = db.query(PropuestaCompra).filter(PropuestaCompra.id == propuesta_id).first()
    if propuesta is None:
        raise ValueError("Propuesta no encontrada")
    if not propuesta.items:
        raise ValueError("No se puede cerrar una propuesta sin ítems")
    propuesta.estado = "cerrada"
    db.commit()
    return obtener_detalle(db, propuesta_id)


def generar_excel(db: Session, propuesta_id: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    detalle = obtener_detalle(db, propuesta_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Propuesta de compra"

    ws.append([f"Propuesta de compra #{detalle['propuesta_id']}"])
    ws.append([f"Responsable: {detalle['responsable'] or '-'}"])
    ws.append([f"Fecha: {detalle['fecha'].strftime('%Y-%m-%d %H:%M')}"])
    ws.append([f"Estado: {detalle['estado']}"])
    ws.append([])

    encabezados = ["Código", "Tela", "% Reposición", "Metros sugeridos", "Metros ajustados",
                   "Metros finales", "Rollos", "Motivo del ajuste"]
    ws.append(encabezados)
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)

    for item in detalle["items"]:
        ws.append([
            item["tela_codigo"], item["tela_nombre"],
            item["pct_reposicion"], item["metros_sugeridos"], item["metros_ajustados"],
            item["metros_finales"], item["rollos"], item["motivo_ajuste"] or "",
        ])

    ws.append([])
    ws.append(["TOTAL", "", "", "", "", detalle["total_metros"], detalle["total_rollos"], ""])
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)

    for col in ws.columns:
        longitud = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(longitud + 2, 10), 40)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_pdf(db: Session, propuesta_id: int) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    detalle = obtener_detalle(db, propuesta_id)
    styles = getSampleStyleSheet()

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    elementos = [
        Paragraph(f"Propuesta de compra #{detalle['propuesta_id']}", styles["Title"]),
        Paragraph(f"Responsable: {detalle['responsable'] or '-'}", styles["Normal"]),
        Paragraph(f"Fecha: {detalle['fecha'].strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Paragraph(f"Estado: {detalle['estado']}", styles["Normal"]),
        Paragraph("Esta es una propuesta de compra, no una orden de compra.", styles["Italic"]),
        Spacer(1, 0.6 * cm),
    ]

    datos_tabla = [["Código", "Tela", "% Repos.", "Metros sug.", "Metros finales", "Rollos"]]
    for item in detalle["items"]:
        datos_tabla.append([
            item["tela_codigo"],
            item["tela_nombre"],
            f"{item['pct_reposicion']:.1f}%" if item["pct_reposicion"] is not None else "-",
            f"{item['metros_sugeridos']:.1f}" if item["metros_sugeridos"] is not None else "-",
            f"{item['metros_finales']:.1f}" if item["metros_finales"] is not None else "-",
            str(item["rollos"] or 0),
        ])
    datos_tabla.append(["", "", "", "", "TOTAL", ""])
    datos_tabla.append(["", "", "", "", f"{detalle['total_metros']:.1f} m", f"{detalle['total_rollos']} rollos"])

    tabla = Table(datos_tabla, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -3), 0.5, colors.grey),
                ("FONTNAME", (-2, -2), (-1, -1), "Helvetica-Bold"),
                ("ALIGN", (2, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    elementos.append(tabla)
    doc.build(elementos)
    return buffer.getvalue()
